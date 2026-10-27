"""
db_export_excel.py – Dienstplan als Excel-Übersicht
====================================================
Exportiert alle MA und ihre Dienste eines Monats als formatierte
Excel-Tabelle: Zeile = MA, Spalte = Tag (01–31)

Aufruf:
  python db_export_excel.py --month März --year 2026
  python db_export_excel.py --month März --year 2026 --out uebersicht.xlsx
"""

import argparse
import sqlite3
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, GradientFill, PatternFill, Side
)
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent.parent))
from config import DB_PATH, MONATE_DE

# ── Farben ────────────────────────────────────────────────────────────────────
COL_HEADER_BG   = "2C3E50"   # dunkelblau – Spaltenheader (Tage)
COL_HEADER_FG   = "FFFFFF"
ROW_HEADER_BG   = "ECF0F1"   # hellgrau – MA-Namen
ROW_HEADER_FG   = "1A1A2E"
WE_COL_BG       = "D5E8F8"   # hellblau – Wochenende-Spalten
WE_COL_DARK     = "B8D4EE"
HEUTE_BG        = "FFF3CD"   # gelb – heutiger Tag
EMPTY_BG        = "FAFAFA"
ALT_ROW_BG      = "F2F2F2"

# Dienst-Kategorie → Hintergrundfarbe in der Tabelle
KATEGORIE_FARBEN = {
    "Arbeit":       None,       # Standard (weiss/grau)
    "Frei":         "E8F8E8",   # hellgrün
    "Abwesenheit":  "FDE8D8",   # lachs
    "Sonstiges":    "F5F5F5",   # grau
}


def get_weekday_short(year: int, month: int, day: int) -> str:
    try:
        return ["Mo","Di","Mi","Do","Fr","Sa","So"][date(year, month, day).weekday()]
    except ValueError:
        return ""


def is_weekend(year: int, month: int, day: int) -> bool:
    try:
        return date(year, month, day).weekday() >= 5
    except ValueError:
        return False


def days_in_month(year: int, month: int) -> int:
    import calendar
    return calendar.monthrange(year, month)[1]


def export(month: str, year: int, out_path: Path) -> None:
    monat_num = MONATE_DE.get(month)
    if not monat_num:
        print(f"❌ Unbekannter Monat: '{month}'")
        sys.exit(1)

    if not DB_PATH.exists():
        print(f"❌ Datenbank nicht gefunden: {DB_PATH}")
        sys.exit(1)

    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    # Alle MA mit Schichten in diesem Monat laden
    rows = conn.execute("""
        SELECT
            m.nachname, m.vorname, m.kuerzel, m.pensum_pct,
            s.datum, s.anzeigename, s.dienst_code,
            COALESCE(dt.kategorie, 'Arbeit') as kategorie
        FROM schichten s
        JOIN mitarbeiter m ON s.mitarbeiter_id = m.id
        LEFT JOIN dienst_typen dt ON s.dienst_code = dt.code
        WHERE strftime('%m', s.datum) = ?
          AND strftime('%Y', s.datum) = ?
          AND m.aktiv = 1
        ORDER BY m.nachname, m.vorname, s.datum
    """, (f"{monat_num:02d}", str(year))).fetchall()

    if not rows:
        print(f"⚠️  Keine Daten für {month} {year} in der DB")
        conn.close()
        sys.exit(0)

    # Daten strukturieren: { kuerzel → { tag → (anzeigename, kategorie) } }
    from collections import defaultdict, OrderedDict
    ma_data: dict = OrderedDict()
    ma_meta: dict = {}

    for r in rows:
        key = r["kuerzel"]   # kuerzel ist eindeutig, nachname/vorname können NULL sein
        if key not in ma_data:
            ma_data[key] = {}
            ma_meta[key] = {"kuerzel": r["kuerzel"], "pensum": r["pensum_pct"], "vorname": r["vorname"], "nachname": r["nachname"]}
        day = int(r["datum"].split("-")[2])
        ma_data[key][day] = (r["anzeigename"] or r["dienst_code"] or "", r["kategorie"])

    n_days = days_in_month(year, monat_num)
    today  = date.today()

    # ── Workbook aufbauen ──────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = f"{month} {year}"

    # Hilfsfunktionen für Styles
    def header_font(bold=True):
        return Font(name="Arial", bold=bold, size=9, color=COL_HEADER_FG)

    def cell_font(bold=False, size=9):
        return Font(name="Arial", bold=bold, size=size)

    def thin_border():
        s = Side(style="thin", color="CCCCCC")
        return Border(left=s, right=s, top=s, bottom=s)

    def fill(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    def center():
        return Alignment(horizontal="center", vertical="center", wrap_text=False)

    def left_align():
        return Alignment(horizontal="left", vertical="center", wrap_text=False)

    # ── Zeile 1: Titel ─────────────────────────────────────────────────────────
    ws.merge_cells(f"A1:{get_column_letter(3 + n_days)}1")
    title_cell = ws["A1"]
    title_cell.value = f"Dienstplan {month} {year}"
    title_cell.font  = Font(name="Arial", bold=True, size=13, color="2C3E50")
    title_cell.alignment = left_align()
    ws.row_dimensions[1].height = 22

    # ── Zeile 2: Spalten-Header (Name | Pensum | 01 Mo | 02 Di | ...) ─────────
    HEADER_ROW = 2
    ws.cell(HEADER_ROW, 1, "Name").font      = Font(name="Arial", bold=True, size=9, color=ROW_HEADER_FG)
    ws.cell(HEADER_ROW, 2, "Kürzel").font    = Font(name="Arial", bold=True, size=9, color=ROW_HEADER_FG)
    ws.cell(HEADER_ROW, 3, "%").font         = Font(name="Arial", bold=True, size=9, color=ROW_HEADER_FG)
    for col in range(1, 4):
        c = ws.cell(HEADER_ROW, col)
        c.fill      = fill(ROW_HEADER_BG)
        c.alignment = center()
        c.border    = thin_border()

    for day in range(1, n_days + 1):
        col     = day + 3
        wd      = get_weekday_short(year, monat_num, day)
        is_we   = is_weekend(year, monat_num, day)
        is_tod  = (today == date(year, monat_num, day))

        c = ws.cell(HEADER_ROW, col, f"{day:02d}\n{wd}")
        c.font      = Font(name="Arial", bold=True, size=8,
                           color="FFFFFF" if not is_we else "1A4A8A")
        c.fill      = fill(WE_COL_DARK if is_we else COL_HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = thin_border()

    ws.row_dimensions[HEADER_ROW].height = 28

    # ── Datenzeilen ────────────────────────────────────────────────────────────
    for ma_idx, (key, day_map) in enumerate(ma_data.items()):
        meta    = ma_meta[key]
        row_num = HEADER_ROW + 1 + ma_idx
        is_alt  = (ma_idx % 2 == 1)
        row_bg  = ALT_ROW_BG if is_alt else "FFFFFF"

        # key = kuerzel (z.B. "von allmen mara")
        vorname  = meta.get("vorname")  or ""
        nachname = meta.get("nachname") or ""
        if nachname or vorname:
            name_str = f"{nachname} {vorname}".strip()
        else:
            name_str = key
        c = ws.cell(row_num, 1, name_str)
        c.font      = Font(name="Arial", bold=False, size=9)
        c.fill      = fill(ROW_HEADER_BG)
        c.alignment = left_align()
        c.border    = thin_border()

        # Kürzel
        c = ws.cell(row_num, 2, meta["kuerzel"] or "")
        c.font      = Font(name="Arial", size=8, color="666666")
        c.fill      = fill(ROW_HEADER_BG)
        c.alignment = center()
        c.border    = thin_border()

        # Pensum
        c = ws.cell(row_num, 3, f"{meta['pensum'] or ''}%")
        c.font      = Font(name="Arial", size=8, color="666666")
        c.fill      = fill(ROW_HEADER_BG)
        c.alignment = center()
        c.border    = thin_border()

        # Tage
        for day in range(1, n_days + 1):
            col    = day + 3
            is_we  = is_weekend(year, monat_num, day)
            is_tod = (today == date(year, monat_num, day))

            if day in day_map:
                label, kategorie = day_map[day]
                kat_bg = KATEGORIE_FARBEN.get(kategorie)
                if is_tod:
                    bg = HEUTE_BG
                elif kat_bg:
                    bg = kat_bg
                elif is_we:
                    bg = WE_COL_BG
                else:
                    bg = row_bg

                c = ws.cell(row_num, col, label)
                c.font      = Font(name="Arial", size=8,
                                   bold=(kategorie == "Arbeit"),
                                   color="1A1A2E")
            else:
                bg = WE_COL_BG if is_we else row_bg
                c  = ws.cell(row_num, col, "")

            if is_tod:
                bg = HEUTE_BG
            elif is_we and day not in day_map:
                bg = WE_COL_BG

            c.fill      = fill(bg)
            c.alignment = center()
            c.border    = thin_border()

        ws.row_dimensions[row_num].height = 15

    # ── Spaltenbreiten ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 20   # Name
    ws.column_dimensions["B"].width = 8    # Kürzel
    ws.column_dimensions["C"].width = 5    # Pensum
    for day in range(1, n_days + 1):
        ws.column_dimensions[get_column_letter(day + 3)].width = 6.5

    # ── Fenster fixieren (Name + Kürzel + % einfrieren) ───────────────────────
    ws.freeze_panes = "D2"

    # ── Legende ───────────────────────────────────────────────────────────────
    legend_row = HEADER_ROW + len(ma_data) + 2
    ws.cell(legend_row, 1, "Legende:").font = Font(name="Arial", bold=True, size=8)
    items = [
        ("Arbeit", "FFFFFF", "Dienste (weiss)"),
        ("Frei", "E8F8E8", "Frei / Urlaub"),
        ("Abwesenheit", "FDE8D8", "Abwesenheit (Krankheit etc.)"),
        ("Wochenende", WE_COL_BG, "Wochenende"),
        ("Heute", HEUTE_BG, "Heutiger Tag"),
    ]
    for i, (label, bg, desc) in enumerate(items):
        col = 2 + i * 2
        c = ws.cell(legend_row, col, label)
        c.fill      = fill(bg)
        c.font      = Font(name="Arial", size=8)
        c.border    = thin_border()
        c.alignment = center()
        ws.cell(legend_row, col + 1, desc).font = Font(name="Arial", size=8, color="666666")

    # ── Speichern ──────────────────────────────────────────────────────────────
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))

    n_ma     = len(ma_data)
    n_eintr  = sum(len(v) for v in ma_data.values())
    print(f"✅ Excel erstellt: {out_path}")
    print(f"   {n_ma} Mitarbeiter, {n_eintr} Einträge, {n_days} Tage")

    conn.close()


def main():
    parser = argparse.ArgumentParser(description="Dienstplan → Excel-Übersicht")
    parser.add_argument("--month", required=True, help="Monat (z.B. März)")
    parser.add_argument("--year",  required=True, type=int, help="Jahr (z.B. 2026)")
    parser.add_argument("--out",   help="Ausgabepfad (Standard: excel_export/Uebersicht_<Monat>_<Jahr>.xlsx)")
    args = parser.parse_args()

    out_path = Path(args.out) if args.out else \
        Path(__file__).parent / "excel_export" / f"Uebersicht_{args.month}_{args.year}.xlsx"

    export(args.month, args.year, out_path)


if __name__ == "__main__":
    main()
